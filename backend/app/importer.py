"""Validate an XLSX template and import it as one CLI operation."""

import re
import uuid
from collections import defaultdict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import TypeAdapter, ValidationError
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from sqlalchemy.orm import Session

from app.auth.models import User
from app.auth.service import normalize_username
from app.catalog import service as catalog_service
from app.catalog.models import Product
from app.catalog.normalize import normalize_article
from app.catalog.schemas import ProductCreate
from app.errors import AppError, validation_message
from app.receipts import service as receipt_service
from app.receipts.schemas import ReceiptCreate, ReceiptLineIn, UnitCost
from app.schema_types import MAX_LINES, Qty

HEADERS = {
    "article": "Артикул",
    "name": "Название",
    "unit": "Единица",
    "sale_price": "Цена продажи, ₸",
    "unit_cost": "Закупочная цена, ₸",
    "qty": "Остаток",
    "min_qty": "Мин. остаток",
    "note": "Примечание",
}
REQUIRED_HEADERS = tuple(value for key, value in HEADERS.items() if key != "note")
UNITS = ("шт", "компл.", "пара", "л", "кг", "м", "упак.")
_QTY = TypeAdapter(Qty)
_UNIT_COST = TypeAdapter(UnitCost)


@dataclass(frozen=True)
class ImportIssue:
    row: int | None
    column: str
    message: str

    def __str__(self) -> str:
        location = f"Строка {self.row}, колонка «{self.column}»" if self.row else self.column
        return f"{location}: {self.message}"


class ImportFailure(Exception):
    """CLI-only errors; they do not extend the public API error codes."""

    def __init__(self, issues: list[ImportIssue]):
        self.issues = issues
        super().__init__("\n".join(map(str, issues)))


@dataclass(frozen=True)
class ImportRow:
    row: int
    product: ProductCreate
    qty: int
    unit_cost: int | None


@dataclass(frozen=True)
class ImportSummary:
    products: int
    with_stock: int
    total_qty: int
    cost_value: int
    retail_value: int


@dataclass(frozen=True)
class ImportResult:
    summary: ImportSummary
    warnings: list[ImportIssue]
    receipt_numbers: list[int]


def _empty(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _integer(
    value: object,
    row: int,
    column: str,
    issues: list[ImportIssue],
    *,
    default: int | None = None,
    required: bool = False,
) -> int | None:
    if _empty(value):
        if required:
            issues.append(ImportIssue(row, column, "Обязательное поле"))
        return default
    try:
        if isinstance(value, bool) or not isinstance(value, int | float | str | Decimal):
            raise ValueError
        if isinstance(value, str) and not re.fullmatch(r"[+-]?[0-9]+", value.strip()):
            raise ValueError
        # XLSX numeric cells may arrive as floats. Convert before any money arithmetic.
        number = Decimal(str(value).strip())
        if not number.is_finite() or number != number.to_integral_value():
            raise ValueError
        if number < 0:
            issues.append(ImportIssue(row, column, "Должно быть не меньше 0"))
            return None
        return int(number)
    except (InvalidOperation, ValueError):
        issues.append(ImportIssue(row, column, "Должно быть целым числом"))
        return None


def _article_text(value: object) -> object:
    if isinstance(value, int | float | Decimal) and not isinstance(value, bool):
        number = Decimal(str(value))
        if number.is_finite():
            result = format(number, "f")
            return result.rstrip("0").rstrip(".") if "." in result else result
    return "" if value is None else value


def _schema_issues(
    exc: ValidationError, row: int, issues: list[ImportIssue], column: str | None = None
) -> None:
    for error in exc.errors():
        header = column or HEADERS[str(error["loc"][0])]
        if not any(issue.row == row and issue.column == header for issue in issues):
            message = validation_message(error)
            if (
                header in (HEADERS["sale_price"], HEADERS["unit_cost"])
                and error["type"] == "less_than_equal"
            ):
                message = f"Должно быть не больше {error['ctx']['le'] // 100} тенге"
            issues.append(ImportIssue(row, header, message))


def _parse_row(
    values: dict[str, object],
    row: int,
    issues: list[ImportIssue],
    article_rows: dict[str, list[int]],
) -> ImportRow | None:
    before = len(issues)
    article = _article_text(values.get("Артикул"))
    if isinstance(article, str) and article.strip():
        try:
            article_rows[normalize_article(article)].append(row)
        except AppError as exc:
            issues.append(ImportIssue(row, "Артикул", exc.message))

    sale_price = _integer(
        values.get(HEADERS["sale_price"]), row, HEADERS["sale_price"], issues, required=True
    )
    unit_cost = _integer(values.get(HEADERS["unit_cost"]), row, HEADERS["unit_cost"], issues)
    qty = _integer(values.get("Остаток"), row, "Остаток", issues, default=0)
    min_qty = _integer(values.get("Мин. остаток"), row, "Мин. остаток", issues)
    if qty is not None and qty > 0:
        if _empty(values.get(HEADERS["unit_cost"])):
            issues.append(
                ImportIssue(row, HEADERS["unit_cost"], "Обязательна при остатке больше 0")
            )
        try:
            _QTY.validate_python(qty)
        except ValidationError as exc:
            _schema_issues(exc, row, issues, "Остаток")
    if unit_cost is not None:
        unit_cost *= 100
        try:
            _UNIT_COST.validate_python(unit_cost)
        except ValidationError as exc:
            _schema_issues(exc, row, issues, HEADERS["unit_cost"])

    unit = values.get("Единица")
    unit = "шт" if _empty(unit) else unit.strip() if isinstance(unit, str) else unit
    if unit not in UNITS:
        issues.append(
            ImportIssue(row, "Единица", f"Неизвестная единица; допустимы: {', '.join(UNITS)}")
        )
    try:
        product = ProductCreate(
            article=article,
            name=values.get("Название") or "",
            unit=unit,
            sale_price=(sale_price or 0) * 100,
            min_qty=min_qty,
            note=values.get("Примечание"),
        )
    except ValidationError as exc:
        _schema_issues(exc, row, issues)
        return None
    if len(issues) != before or qty is None:
        return None
    return ImportRow(row, product, qty, unit_cost)


def _read_rows(
    path: Path,
) -> tuple[list[ImportRow], dict[str, list[int]], list[ImportIssue]]:
    issues: list[ImportIssue] = []
    rows: list[ImportRow] = []
    article_rows: dict[str, list[int]] = defaultdict(list)
    book = None
    try:
        book = load_workbook(path, read_only=True, data_only=False)
        if "Товары" not in book.sheetnames:
            return rows, article_rows, [ImportIssue(None, "Файл", "Нет листа «Товары»")]
        source = book["Товары"].iter_rows()
        first = next(source, ())
        columns: dict[str, int] = {}
        for index, cell in enumerate(first):
            if cell.value in HEADERS.values():
                if cell.value in columns:
                    issues.append(ImportIssue(1, cell.value, "Заголовок повторяется"))
                columns[cell.value] = index
        for header in REQUIRED_HEADERS:
            if header not in columns:
                issues.append(ImportIssue(1, header, "Нет обязательного заголовка"))
        if issues:
            return rows, article_rows, issues
        for row_number, cells in enumerate(source, start=2):
            values = {header: cells[index].value for header, index in columns.items()}
            if all(_empty(value) for value in values.values()):
                continue
            formulas = [
                header for header, index in columns.items() if cells[index].data_type == "f"
            ]
            for header in formulas:
                issues.append(
                    ImportIssue(row_number, header, "Формулы не поддерживаются; укажите значение")
                )
                values[header] = None
            parsed = _parse_row(values, row_number, issues, article_rows)
            if parsed is not None and not formulas:
                rows.append(parsed)
    except (OSError, BadZipFile, InvalidFileException, ParseError, KeyError, ValueError) as exc:
        raise ImportFailure([ImportIssue(None, "Файл", "Не удалось прочитать XLSX")]) from exc
    finally:
        if book is not None:
            book.close()
    return rows, article_rows, issues


def _existing_issues(db: Session, article_rows: dict[str, list[int]]) -> list[ImportIssue]:
    existing = db.scalars(
        select(Product.article_norm).where(Product.article_norm.in_(article_rows))
    )
    return [
        ImportIssue(row, "Артикул", "Артикул уже есть в базе")
        for norm in existing
        for row in article_rows[norm]
    ]


def import_products(db: Session, path: Path, username: str, *, apply: bool = False) -> ImportResult:
    """Validate every row first; only this top-level operation may commit the import."""
    rows, article_rows, issues = _read_rows(path)
    user_id = db.scalar(select(User.id).where(User.username == normalize_username(username)))
    if user_id is None:
        issues.append(ImportIssue(None, "Параметр --user", "Пользователь не найден"))
    for numbers in article_rows.values():
        if len(numbers) > 1:
            detail = f"Повтор артикула после нормализации в строках {', '.join(map(str, numbers))}"
            issues.extend(ImportIssue(row, "Артикул", detail) for row in numbers)
    issues.extend(_existing_issues(db, article_rows))
    if issues:
        raise ImportFailure(sorted(issues, key=lambda issue: (issue.row or 0, issue.column)))

    summary = ImportSummary(
        products=len(rows),
        with_stock=sum(row.qty > 0 for row in rows),
        total_qty=sum(row.qty for row in rows),
        cost_value=sum(row.qty * (row.unit_cost or 0) for row in rows),
        retail_value=sum(row.qty * row.product.sale_price for row in rows),
    )
    warnings = []
    for row in rows:
        if row.unit_cost is not None and row.unit_cost >= row.product.sale_price:
            warnings.append(
                ImportIssue(row.row, HEADERS["unit_cost"], "Закупочная цена не меньше цены продажи")
            )
        if row.qty == 0:
            warnings.append(ImportIssue(row.row, "Остаток", "Товар с нулевым остатком"))
    numbers: list[int] = []
    if apply:
        current_row: int | None = None
        try:
            lines = []
            for row in rows:
                current_row = row.row
                product = catalog_service._create_product(db, row.product)
                db.flush()
                if row.qty > 0:
                    lines.append(
                        ReceiptLineIn(
                            product_id=product.id, qty=row.qty, unit_price_minor=row.unit_cost
                        )
                    )
            current_row = None
            for start in range(0, len(lines), MAX_LINES):
                receipt = receipt_service._post_receipt(
                    db,
                    ReceiptCreate(
                        request_id=uuid.uuid4(),
                        lines=lines[start : start + MAX_LINES],
                        supplier="X2",
                        note="Бастапқы қалдық, X2-ден көшіру",
                    ),
                    user_id,
                )
                numbers.append(receipt.number)
            db.commit()
        except (AppError, SQLAlchemyError) as exc:
            db.rollback()
            if isinstance(exc, IntegrityError):
                constraint = getattr(getattr(exc.orig, "diag", None), "constraint_name", None)
                if constraint == catalog_service.ARTICLE_NORM_CONSTRAINT:
                    duplicates = _existing_issues(db, article_rows)
                    if duplicates:
                        raise ImportFailure(duplicates) from exc
            detail = exc.message if isinstance(exc, AppError) else "Ошибка записи в базу"
            raise ImportFailure(
                [
                    ImportIssue(
                        current_row,
                        "Артикул" if current_row else "Импорт",
                        f"{detail}. Все изменения импорта отменены",
                    )
                ]
            ) from exc
        except BaseException:
            db.rollback()
            raise
    return ImportResult(summary, warnings, numbers)
